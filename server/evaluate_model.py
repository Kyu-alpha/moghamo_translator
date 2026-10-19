"""
Baseline evaluation of your English -> Meta' model on held-out sentences.
---------------------------------------------------------------------------
Reads the spreadsheet (columns "english" and "meta"), translates every
English sentence with your trained model, and compares the model's output
to the book's Meta'. Nothing here trains the model - it only measures it.

Run it from the server folder, so it can find model_interface.py and the
model/ folder:

    pip install openpyxl
    pip install sacrebleu        (optional - adds a chrF score)
    python evaluate_model.py path\\to\\trouble_in_the_market.xlsx

Output:
  - a summary printed to the screen
  - evaluation_results.csv (open it in Excel) with one row per sentence

HOW TO READ THE SCORES
  The similarity score is character overlap between the model's output and
  the book's text (0 = nothing in common, 1 = identical). It is NOT a
  measure of meaning. A low score can still be a correct translation worded
  differently, and a high score can hide one wrong word. Use the numbers to
  find the rows worth showing to a speaker, then let a speaker judge them.

  Two versions of the score are reported:
    - strict:   tone marks count
    - no tone:  tone marks are ignored (useful if your NT training text
                marks tone differently from this book)

SENTENCE FRAGMENTS
  Rows are used exactly as they are in the spreadsheet. If it still has
  rows that are pieces of one longer sentence, you can merge them here
  without editing the file, e.g.  --merge 13-14,31-33  (spreadsheet row
  numbers, before any merging). By default nothing is merged.
"""

import argparse
import csv
import difflib
import os
import sys
import unicodedata

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)

PUNCT = ".,!?:;\"“”"


def norm(text):
    """Unicode-normalize and tidy whitespace and apostrophes."""
    text = unicodedata.normalize("NFC", text or "")
    for ch in ("’", "‘", "ʼ"):
        text = text.replace(ch, "'")
    return " ".join(text.split())


def strip_tones(text):
    """Remove combining accent marks (tone marks), keep the base letters."""
    decomposed = unicodedata.normalize("NFD", text)
    kept = "".join(ch for ch in decomposed if unicodedata.category(ch) != "Mn")
    return unicodedata.normalize("NFC", kept)


def for_compare(text, tones=True):
    text = norm(text).casefold()
    text = "".join(ch for ch in text if ch not in PUNCT)
    if not tones:
        text = strip_tones(text)
    return " ".join(text.split())


def similarity(a, b, tones=True):
    a, b = for_compare(a, tones), for_compare(b, tones)
    if not a and not b:
        return 1.0
    return difflib.SequenceMatcher(None, a, b).ratio()


def load_pairs(path):
    """Returns a list of (row_number, english, meta), skipping incomplete rows."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb["data"] if "data" in wb.sheetnames else wb.active

    header = {}
    for col, cell in enumerate(ws[1], start=1):
        if cell.value:
            header[str(cell.value).strip().lower()] = col
    if "english" not in header or "meta" not in header:
        sys.exit("Could not find 'english' and 'meta' column headers in row 1 of the sheet.")

    pairs, skipped = [], []
    for row in range(2, ws.max_row + 1):
        en = ws.cell(row=row, column=header["english"]).value
        mg = ws.cell(row=row, column=header["meta"]).value
        if en and mg and str(en).strip() and str(mg).strip():
            pairs.append((row, norm(str(en)), norm(str(mg))))
        elif en or mg:
            skipped.append(row)
    if skipped:
        print(f"Skipped incomplete rows: {skipped}")
    return pairs


def parse_merge(spec):
    if not spec or spec.strip().lower() == "none":
        return []
    ranges = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            ranges.append((int(a), int(b)))
        else:
            ranges.append((int(part), int(part)))
    return ranges


def merge_pairs(pairs, ranges):
    """Joins consecutive rows inside each range into one sentence pair."""
    def find_range(row):
        for a, b in ranges:
            if a <= row <= b:
                return (a, b)
        return None

    out, i = [], 0
    while i < len(pairs):
        row, en, mg = pairs[i]
        rng = find_range(row)
        if rng is None or rng[0] == rng[1]:
            out.append((str(row), en, mg))
            i += 1
            continue
        group = [p for p in pairs if rng[0] <= p[0] <= rng[1]]
        out.append((f"{rng[0]}-{rng[1]}",
                    " ".join(p[1] for p in group),
                    " ".join(p[2] for p in group)))
        i += len(group)
    return out


def encode_target(tok, text):
    try:
        return tok(text_target=text)["input_ids"]
    except Exception:
        return tok(text)["input_ids"]


def tokenizer_report(tok, texts):
    """Checks whether the model's tokenizer can represent the book's Meta' text."""
    unk = getattr(tok, "unk_token_id", None)
    bad_chars = []
    for ch in sorted(set("".join(texts))):
        if ch.isspace():
            continue
        ids = encode_target(tok, ch)
        if unk is not None and unk in ids:
            bad_chars.append(ch)

    mismatches = []
    for text in texts:
        ids = encode_target(tok, text)
        back = tok.decode(ids, skip_special_tokens=True)
        if for_compare(back) != for_compare(text):
            mismatches.append((text, back))
    return bad_chars, mismatches


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="path to trouble_in_the_market.xlsx")
    ap.add_argument("--merge", default="none",
                    help="spreadsheet row ranges to merge into one sentence, e.g. 13-14,31-33 (default: none)")
    ap.add_argument("--out", default="evaluation_results.csv", help="where to write the CSV")
    args = ap.parse_args()

    pairs = merge_pairs(load_pairs(args.xlsx), parse_merge(args.merge))
    if not pairs:
        sys.exit("No complete english/meta rows found.")
    print(f"Loaded {len(pairs)} sentence pairs.\n")

    import model_interface as mi

    print("Loading the model (first run is slow)...")
    mi._load_model()
    tok = mi._tokenizer

    # --- 1. Can the tokenizer even represent this spelling? -------------------
    print("\nTokenizer check on the book's Meta' text")
    try:
        bad_chars, mismatches = tokenizer_report(tok, [p[2] for p in pairs])
        if bad_chars:
            print("  Characters the tokenizer maps to UNKNOWN:", " ".join(bad_chars))
            print("  -> The model has never seen these in training. Fix the spelling or retrain with them.")
        else:
            print("  No unknown characters.")
        print(f"  Text that does not survive encode -> decode unchanged: {len(mismatches)} of {len(pairs)}")
        for orig, back in mismatches[:3]:
            print(f"    book:    {orig}\n    decoded: {back}")
    except Exception as exc:  # tokenizer APIs differ between models
        print(f"  (skipped: {exc})")

    # --- 2. Translate every English sentence ----------------------------------
    print("\nTranslating...")
    results = []
    for n, (label, en, book) in enumerate(pairs, start=1):
        try:
            model_out = mi.translate_en_to_moghamo(en)
        except Exception as exc:
            print(f"  row {label}: model error: {exc}")
            model_out = ""
        sim = similarity(model_out, book, tones=True)
        sim_nt = similarity(model_out, book, tones=False)
        exact = for_compare(model_out) == for_compare(book)
        exact_nt = for_compare(model_out, False) == for_compare(book, False)
        results.append((label, en, book, model_out, sim, sim_nt, exact, exact_nt))
        print(f"  [{n}/{len(pairs)}] row {label}  similarity {sim:.2f}", flush=True)

    # --- 3. Write the CSV -----------------------------------------------------
    with open(args.out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["sheet_row", "english", "book_meta", "model_meta",
                    "similarity", "similarity_no_tone", "exact", "exact_no_tone"])
        for label, en, book, out, s, snt, ex, exnt in results:
            w.writerow([label, en, book, out, f"{s:.3f}", f"{snt:.3f}", ex, exnt])

    # --- 4. Summary -----------------------------------------------------------
    n = len(results)
    mean = lambda idx: sum(r[idx] for r in results) / n
    print("\n" + "=" * 60)
    print(f"Sentences evaluated:            {n}")
    print(f"Mean similarity (strict):       {mean(4):.3f}")
    print(f"Mean similarity (no tone):      {mean(5):.3f}")
    print(f"Exact matches (strict):         {sum(r[6] for r in results)} / {n}")
    print(f"Exact matches (no tone):        {sum(r[7] for r in results)} / {n}")
    high = sum(1 for r in results if r[4] >= 0.8)
    mid = sum(1 for r in results if 0.5 <= r[4] < 0.8)
    low = sum(1 for r in results if r[4] < 0.5)
    print(f"Similarity >= 0.8 / 0.5-0.8 / < 0.5:  {high} / {mid} / {low}")

    try:
        from sacrebleu.metrics import CHRF
        refs = [for_compare(r[2]) for r in results]
        hyps = [for_compare(r[3]) for r in results]
        refs_nt = [for_compare(r[2], False) for r in results]
        hyps_nt = [for_compare(r[3], False) for r in results]
        print(f"chrF (strict / no tone):        {CHRF().corpus_score(hyps, [refs]).score:.1f} / "
              f"{CHRF().corpus_score(hyps_nt, [refs_nt]).score:.1f}")
    except ImportError:
        print("chrF: not installed (pip install sacrebleu to add it)")

    print("\nLowest-scoring rows (show these to a speaker first):")
    for label, en, book, out, s, snt, ex, exnt in sorted(results, key=lambda r: r[4])[:5]:
        print(f"  row {label}  similarity {s:.2f}\n    English: {en}\n    Book:    {book}\n    Model:   {out}")

    print(f"\nFull results saved to {args.out}")
    print("Remember: similarity measures spelling overlap, not meaning. A speaker has the final say.")


if __name__ == "__main__":
    main()
